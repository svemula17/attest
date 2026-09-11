"""Questionnaire import (CSV / XLSX, problems with row numbers) and export (CSV / XLSX), plus the CLI."""
from __future__ import annotations

import argparse
import csv
import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from attest.auth import SYSTEM
from attest.config import default_config_toml, parse_config
from attest.db import get_engine, upgrade
from attest.questionnaires import EXPORT_COLUMNS, NOT_ANSWERED, QuestionnaireError, export_questionnaire, parse_questionnaire
from attest.service import Service
from attest.store_sql import QuestionnaireStore


@pytest.fixture
def sandbox(tmp_path):
    cfg_path = tmp_path / "attest.toml"
    cfg_path.write_text(default_config_toml(mode="sandbox", storage_url=f"sqlite:///{tmp_path / 'attest.db'}"))
    cfg = parse_config(cfg_path.read_text(), path=cfg_path)
    upgrade(cfg.storage.url)
    engine = get_engine(cfg.storage.url)
    svc = Service(cfg, engine)
    svc.reseed()
    yield svc
    engine.dispose()


def _csv(path: Path, text: str) -> Path:
    path.write_text(text, encoding="utf-8")
    return path


def _xlsx(path: Path, rows: list[list], second_sheet: list[list] | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Questionnaire"
    for row in rows:
        ws.append(row)
    if second_sheet:
        other = wb.create_sheet("Ignored")
        for row in second_sheet:
            other.append(row)
    wb.save(path)
    return path


def _read_csv(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _read_xlsx(path: Path) -> list[dict]:
    ws = load_workbook(io.BytesIO(path.read_bytes()), read_only=True).worksheets[0]  # bytes: openpyxl gates on the path's suffix
    rows = [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:]]


# -- parse: CSV ------------------------------------------------------------------------

def test_parse_csv_columns_are_case_insensitive_and_controls_split_on_either_separator(tmp_path):
    path = _csv(tmp_path / "acme.csv", "﻿ID,Question,Control_IDs\n"
                                        "1.1,Is data encrypted at rest?,CTL-CRYPTO-01; CTL-CRYPTO-02\n"
                                        "1.2,  Do you review privileged access?  ,\"CTL-PRIV-01,CTL-ACCESS-01\"\n"
                                        ",,\n"
                                        "1.3,Any breaches in 24 months?,\n"
                                        "Section 2,,\n")
    rows = parse_questionnaire(path)
    assert rows == [
        {"question": "Is data encrypted at rest?", "control_ids": ["CTL-CRYPTO-01", "CTL-CRYPTO-02"], "ref": "1.1"},
        {"question": "Do you review privileged access?", "control_ids": ["CTL-PRIV-01", "CTL-ACCESS-01"], "ref": "1.2"},
        {"question": "Any breaches in 24 months?", "control_ids": [], "ref": "1.3"},
    ]


def test_parse_csv_question_only(tmp_path):
    path = _csv(tmp_path / "q.csv", "question\nOne?\n\nTwo?\n")
    assert parse_questionnaire(path) == [{"question": "One?", "control_ids": [], "ref": None},
                                         {"question": "Two?", "control_ids": [], "ref": None}]


def test_parse_reports_every_problem_with_row_numbers(tmp_path):
    path = _csv(tmp_path / "bad.csv", "ref,question,control_ids\n"
                                       "A1,Encrypted?,CTL-CRYPTO-01\n"
                                       "A2,Reviewed?,CTL PRIV 01;CTL-ACCESS-01\n"
                                       "A1,Backups?,CTL-BACKUP-01\n"
                                       ",Heading row without a question,\n"
                                       "A4,Logged?,;;\n"
                                       "A5,Long?," + "x" * 4001 + "\n")
    with pytest.raises(QuestionnaireError) as exc:
        parse_questionnaire(path)
    problems = exc.value.problems
    assert problems == [
        "row 3: invalid control id(s) 'CTL PRIV 01'; use ids like CTL-CRYPTO-01 separated by ';'",
        "row 4: duplicate ref 'A1' (first used at row 2)",
    ]
    assert exc.value.path == path and "bad.csv: 2 problems" in str(exc.value) and "row 3:" in str(exc.value)

    path = _csv(tmp_path / "long.csv", "question\n" + "x" * 4001 + "\n")
    with pytest.raises(QuestionnaireError, match=r"row 2: question is 4001 characters long"):
        parse_questionnaire(path)


def test_parse_header_problems(tmp_path):
    with pytest.raises(QuestionnaireError, match=r"row 1: header has no 'question' column \(found: 'id', 'prompt'\)"):
        parse_questionnaire(_csv(tmp_path / "noq.csv", "id,prompt\n1,Hello?\n"))
    with pytest.raises(QuestionnaireError, match="row 1: header names column 'question' twice"):
        parse_questionnaire(_csv(tmp_path / "dup.csv", "question,Question\na,b\n"))
    with pytest.raises(QuestionnaireError, match="the file is empty"):
        parse_questionnaire(_csv(tmp_path / "empty.csv", "\n\n"))
    with pytest.raises(QuestionnaireError, match="no questions found"):
        parse_questionnaire(_csv(tmp_path / "blank.csv", "ref,question\n1,\n2,   \n"))
    (tmp_path / "q.docx").write_bytes(b"PK")
    with pytest.raises(QuestionnaireError, match="unsupported file type '.docx'"):
        parse_questionnaire(tmp_path / "q.docx")
    with pytest.raises(QuestionnaireError, match="no such file"):
        parse_questionnaire(tmp_path / "missing.csv")


# -- parse: XLSX -----------------------------------------------------------------------

def test_parse_xlsx_first_sheet_header_row_and_numeric_refs(tmp_path):
    path = _xlsx(tmp_path / "acme.xlsx", [
        [None, None, None],                                   # leading blank row
        ["Ref", "QUESTION", "control_ids", "notes"],
        [3, "Encrypted at rest?", "CTL-CRYPTO-01", "ignored"],
        [4.0, "Privileged access reviewed?", "CTL-PRIV-01, CTL-ACCESS-01", None],
        [None, None, None, None],
        ["5a", None, "CTL-LOG-01", "blank question: skipped"],
        [6, " Incident SLA? ", None, None],
    ], second_sheet=[["question"], ["Should not be read"]])
    rows = parse_questionnaire(path)
    assert rows == [
        {"question": "Encrypted at rest?", "control_ids": ["CTL-CRYPTO-01"], "ref": "3"},
        {"question": "Privileged access reviewed?", "control_ids": ["CTL-PRIV-01", "CTL-ACCESS-01"], "ref": "4"},
        {"question": "Incident SLA?", "control_ids": [], "ref": "6"},
    ]


def test_parse_xlsx_problems_carry_sheet_row_numbers(tmp_path):
    path = _xlsx(tmp_path / "bad.xlsx", [
        ["question", "control_ids", "id"],
        ["One?", "CTL-A", "r1"],
        ["Two?", "bad id", "r2"],
        ["Three?", "CTL-C", "r1"],
    ])
    with pytest.raises(QuestionnaireError) as exc:
        parse_questionnaire(path)
    assert exc.value.problems == ["row 3: invalid control id(s) 'bad id'; use ids like CTL-CRYPTO-01 separated by ';'",
                                  "row 4: duplicate ref 'r1' (first used at row 2)"]
    (tmp_path / "corrupt.xlsx").write_bytes(b"not a workbook")
    with pytest.raises(QuestionnaireError, match="not a readable .xlsx workbook"):
        parse_questionnaire(tmp_path / "corrupt.xlsx")
    with pytest.raises(QuestionnaireError, match="the file is empty"):
        parse_questionnaire(_xlsx(tmp_path / "empty.xlsx", [[None, None]]))


# -- export ----------------------------------------------------------------------------

ROWS = [
    dict(ref="1.1", question_id="Q-4475", question="Encrypted at rest?", status="READY", answer="Yes, AES-256 under KMS.",
         reason="", citations=[{"id": "EV-0004", "source": "aws-config", "collected_at": "2026-09-11T10:00:00Z"},
                               {"id": "EV-0005", "source": "aws-config", "collected_at": "2026-09-11T10:00:00Z"}],
         decision="approve", approver="s.vemula@attest.internal", signed_at="2026-09-11T12:00:00Z", signature="ab" * 32),
    dict(ref="1.2", question_id="Q-4476", question="Any breaches?", status="DECLINED", answer="",
         reason="no evidence cites the requested controls", citations=[], decision=None, approver=None, signed_at=None, signature=None),
    dict(question_id="Q-4477", question="Everything?", status="BLOCKED", answer="", reason="max-steps: 24 controls requested",
         citations=[], decision=None, approver=None, signed_at=None, signature=None),
    dict(ref="1.4", question_id="Q-4478", question="Pending?", status="READY", answer="Draft answer.", reason="",
         citations=[{"id": "EV-0010"}], decision=None, approver=None, decided_at=None, signature=None),
]

EXPECTED = [
    ["1.1", "Q-4475", "Encrypted at rest?", "READY", "Yes, AES-256 under KMS.", "EV-0004;EV-0005", "approve",
     "s.vemula@attest.internal", "2026-09-11T12:00:00Z", "ab" * 32],
    ["1.2", "Q-4476", "Any breaches?", "DECLINED", NOT_ANSWERED + "no evidence cites the requested controls", "", "", "", "", ""],
    ["", "Q-4477", "Everything?", "BLOCKED", NOT_ANSWERED + "max-steps: 24 controls requested", "", "", "", "", ""],
    ["1.4", "Q-4478", "Pending?", "READY", "Draft answer.", "EV-0010", "", "", "", ""],
]


def test_export_csv_columns_and_not_answered_rows(tmp_path):
    out = export_questionnaire(ROWS, tmp_path / "answers.csv")
    assert out == tmp_path / "answers.csv"
    rows = _read_csv(out)
    assert list(rows[0]) == list(EXPORT_COLUMNS)
    assert [[r[c] for c in EXPORT_COLUMNS] for r in rows] == EXPECTED
    # the export is itself a valid questionnaire: refs and questions roundtrip
    again = parse_questionnaire(out)
    assert [(r["ref"], r["question"]) for r in again] == [("1.1", "Encrypted at rest?"), ("1.2", "Any breaches?"),
                                                         (None, "Everything?"), ("1.4", "Pending?")]


def test_export_xlsx_by_extension_and_by_explicit_format(tmp_path):
    out = export_questionnaire(ROWS, tmp_path / "answers.xlsx")
    rows = _read_xlsx(out)
    assert list(rows[0]) == list(EXPORT_COLUMNS)
    assert [[r[c] for c in EXPORT_COLUMNS] for r in rows] == EXPECTED
    ws = load_workbook(out).worksheets[0]
    assert ws.title == "Answers" and ws.freeze_panes == "A2" and ws["A1"].font.bold

    forced = export_questionnaire(ROWS, tmp_path / "answers.dat", fmt="xlsx")
    assert _read_xlsx(forced) == rows
    as_csv = export_questionnaire(ROWS, tmp_path / "answers.xlsx.bak", fmt="csv")
    assert _read_csv(as_csv)[0]["question_id"] == "Q-4475"
    assert _read_csv(export_questionnaire([], tmp_path / "empty.csv")) == []
    with pytest.raises(ValueError, match="fmt must be"):
        export_questionnaire(ROWS, tmp_path / "x.json", fmt="json")


# -- through the service ------------------------------------------------------------------

def test_import_then_export_roundtrip_through_the_service(sandbox, tmp_path):
    path = _csv(tmp_path / "acme.csv", "ref,question,control_ids\n"
                                        "2.1,Is customer PHI encrypted at rest and in transit?,CTL-CRYPTO-01;CTL-CRYPTO-02\n"
                                        "2.2,Have you had a reportable breach in the last 24 months?,\n"
                                        "2.3,Describe your privileged access review cadence.,CTL-PRIV-01\n")
    res = sandbox.import_questionnaire(SYSTEM, path, name="Acme vendor review")
    qn, drafts = res["questionnaire"], res["drafts"]
    assert qn["id"] == "QN-0001" and qn["row_count"] == 3 and qn["status"] == "open" and qn["source_file"] == "acme.csv"
    assert [d["status"] for d in drafts] == ["READY", "DECLINED", "READY"]
    linked = sandbox.drafts.list_for_questionnaire(qn["id"])
    assert [d["question_id"] for d in linked] == [d["question_id"] for d in drafts]
    assert all(d["questionnaire_id"] == qn["id"] for d in linked)
    for row, d in zip(parse_questionnaire(path), drafts):  # what the CLI does with the customer's refs
        sandbox.drafts.set_ref(d["question_id"], row["ref"])
    assert [d["ref"] for d in sandbox.drafts.list_for_questionnaire(qn["id"])] == ["2.1", "2.2", "2.3"]

    sandbox.decide(sandbox._demo_identity(), drafts[0]["question_id"], "approve")
    out = sandbox.export_questionnaire(SYSTEM, qn["id"], tmp_path / "answers.xlsx")
    rows = _read_xlsx(out)
    assert [r["question_id"] for r in rows] == [d["question_id"] for d in drafts]
    assert rows[0]["decision"] == "approve" and len(rows[0]["signature"]) == 64 and rows[0]["citations"]
    assert rows[1]["answer"].startswith(NOT_ANSWERED) and rows[1]["decision"] == ""
    assert rows[2]["status"] == "READY" and rows[2]["decision"] == ""
    assert QuestionnaireStore(sandbox.engine).get(qn["id"])["status"] == "exported"
    assert any(e.action == "questionnaire.exported" and e.subject == qn["id"] for e in sandbox.audit.query(limit=5))
    with pytest.raises(ValueError, match="unknown questionnaire"):
        sandbox.export_questionnaire(SYSTEM, "QN-0042", tmp_path / "x.csv")


# -- CLI ----------------------------------------------------------------------------------

def _cli(*argv: str) -> int:
    from attest.cli_product import register
    p = argparse.ArgumentParser()
    p.add_argument("--config")
    p.add_argument("--data", type=Path, default=Path("data"))
    sub = p.add_subparsers(dest="cmd", required=True)
    register(sub)
    args = p.parse_args(list(argv))
    return args.fn(args)


def test_cli_questionnaire_import_export_list(sandbox, tmp_path, capsys):
    cfg = str(sandbox.cfg.path)
    assert _cli("--config", cfg, "questionnaire", "list") == 0
    assert "no questionnaires" in capsys.readouterr().out

    path = _xlsx(tmp_path / "northwind.xlsx", [
        ["ref", "question", "control_ids"],
        ["A.1", "Is customer PHI encrypted at rest and in transit?", "CTL-CRYPTO-01;CTL-CRYPTO-02"],
        ["A.2", "Any reportable breach in the last 24 months?", ""],
    ])
    assert _cli("--config", cfg, "questionnaire", "import", str(path)) == 0
    text = capsys.readouterr().out
    assert "QN-0001  northwind  —  2 questions from northwind.xlsx" in text
    assert "A.1      Q-4475   READY" in text and "A.2      Q-4476   DECLINED" in text
    assert "ready 1 · declined 1 · blocked 0" in text
    assert [d["ref"] for d in sandbox.drafts.list_for_questionnaire("QN-0001")] == ["A.1", "A.2"]

    assert _cli("--config", cfg, "questionnaire", "list") == 0
    text = capsys.readouterr().out
    assert "QN-0001  open" in text and "2 rows" in text and "1 undecided" in text and "northwind" in text

    out = tmp_path / "answers.csv"
    assert _cli("--config", cfg, "questionnaire", "export", "QN-0001", "--out", str(out)) == 0
    text = capsys.readouterr().out
    assert f"wrote {out} (2 rows · 1 not answered)" in text and "warning: 1 READY answer not yet approved" in text and "Q-4475" in text
    rows = _read_csv(out)
    assert [r["ref"] for r in rows] == ['A.1', 'A.2']  # refs ride on the drafts; Service._queue_item does not surface them yet
    assert rows[0]["status"] == "READY" and rows[1]["answer"].startswith(NOT_ANSWERED)
    assert QuestionnaireStore(sandbox.engine).get("QN-0001")["status"] == "exported"

    assert _cli("--config", cfg, "questionnaire", "export", "QN-0009", "--out", str(tmp_path / "x.csv")) == 1
    assert "unknown questionnaire QN-0009" in capsys.readouterr().out


def test_cli_questionnaire_import_rejects_a_bad_file_without_creating_anything(sandbox, tmp_path, capsys):
    path = _csv(tmp_path / "bad.csv", "ref,question,control_ids\n1,Ok?,CTL-A\n1,Dup ref?,CTL B\n")
    assert _cli("--config", str(sandbox.cfg.path), "questionnaire", "import", str(path), "--name", "Bad") == 2
    text = capsys.readouterr().out
    assert "bad.csv: 2 problems — nothing imported" in text
    assert "  - row 3: invalid control id(s) 'CTL B'" in text and "  - row 3: duplicate ref '1' (first used at row 2)" in text
    assert QuestionnaireStore(sandbox.engine).list() == []
    assert not any(d.get("questionnaire_id") for d in sandbox.drafts.list())
